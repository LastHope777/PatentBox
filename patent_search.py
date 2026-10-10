"""
Обработка входного документа.

Во входном .docx ищутся таблицы, у которых в первой ячейке написан источник
(«ФИПС», «Платформа» или «WIPO»), а ниже — номера патентов. По каждой такой
таблице в result.docx дописывается таблица результата: шапка, номера граф,
пустая строка для темы и по строке на каждый патент.

Сами сайты обрабатываются в scrapers.py, интерфейс — в main.py.
"""

import os
import traceback
from dataclasses import dataclass, field

from docx import Document

import scrapers
from paths import RESULT_DOCX
from result_picker import ExactResultNotFound

FIPS_ALL_DATABASES = scrapers.FIPS_ALL_DATABASES

TABLE_HEADERS = (
    'Объект исследования, его составные части',
    'Вид и номер охранного документа с двухбуквенным кодом страны/режим действия '
    '(действует/не действует/может быть восстановлен)',
    'Название объекта интеллектуальной собственности',
    'Заявитель (правообладатель)',
    'Классификационные рубрики',
    'Дата приоритета/ дата публикации',
)
NOT_FOUND_TEXT = "Документ с таким номером не найден в выдаче — проверьте вручную"
ERROR_TEXT = "Ошибка при обработке ({}) — проверьте вручную"


@dataclass(frozen=True)
class Source:
    title: str
    aliases: tuple            # что может стоять в первой ячейке входной таблицы (без учёта регистра)
    browser_per_patent: bool  # True — новый браузер на каждый патент, False — один на всю таблицу


SOURCES = {
    "fips": Source("ФИПС", ("фипс", "fips"), browser_per_patent=True),
    "platform": Source("Платформа Роспатента", ("платформа", "роспатент"), browser_per_patent=False),
    "wipo": Source("WIPO", ("wipo", "випо"), browser_per_patent=True),
}


@dataclass
class RunSummary:
    result_path: str
    found: list = field(default_factory=list)      # [Patent]
    not_found: list = field(default_factory=list)  # ["ФИПС: 2848852", ...]
    errors: list = field(default_factory=list)     # ["WIPO: CN103900678 — текст ошибки", ...]

    def message(self):
        lines = [f"Данные сохранены в {os.path.basename(self.result_path)}.",
                 f"Найдено патентов: {len(self.found)}."]
        if self.not_found:
            lines.append(f"\nНе найдены в выдаче ({len(self.not_found)}):")
            lines += [f"  • {item}" for item in self.not_found]
        if self.errors:
            lines.append(f"\nОшибки ({len(self.errors)}), подробности в консоли:")
            lines += [f"  • {item}" for item in self.errors]
        return "\n".join(lines)


# ============================== Входной документ ==============================

def detect_source(cell_text):
    """Ключ источника по тексту первой ячейки таблицы или None."""
    text = " ".join(cell_text.split()).lower()
    for key, source in SOURCES.items():
        if text in source.aliases:
            return key
    return None


def read_jobs(input_path):
    """[(ключ_источника, [номера...]), ...] — в том порядке, в каком таблицы идут в файле."""
    jobs = []
    for table in Document(input_path).tables:
        key = detect_source(table.cell(0, 0).text)
        if key is None:
            continue
        requests = (" ".join(row.cells[0].text.split()) for row in table.rows[1:])
        jobs.append((key, [request for request in requests if request]))  # пустые строки пропускаем
    return jobs


# ============================== Таблица результата ==============================

def create_result_table(document):
    """
    Добавляет в документ таблицу результата:
      строка 0 — заголовки граф;
      строка 1 — номера граф 1..6;
      строка 2 — пустая объединённая на всю ширину строка для темы (заполняется вручную).
    """
    table = document.add_table(rows=3, cols=len(TABLE_HEADERS))
    for col, title in enumerate(TABLE_HEADERS):
        table.cell(0, col).text = title
        table.cell(1, col).text = str(col + 1)
    table.cell(2, 0).merge(table.cell(2, len(TABLE_HEADERS) - 1))
    return table


def add_row(table, number, name="", holder="", mpk="", date=""):
    """Новая строка с графами 2–6 (графа 1 «Объект исследования» заполняется вручную)."""
    cells = table.add_row().cells
    for col, value in enumerate((number, name, holder, mpk, date), start=1):
        cells[col].text = value


def _short_error(exc):
    """Первая строка текста ошибки (у Selenium дальше идёт длинный stacktrace)."""
    text = getattr(exc, "msg", None) or str(exc) or type(exc).__name__
    return text.strip().splitlines()[0][:200]


# ============================== Основной процесс ==============================

def process_document(input_path, driver_path, browser, fips_databases=(), result_path=RESULT_DOCX):
    """
    Обрабатывает входной файл и дописывает таблицы в result_path.
      browser        — "google" или "edge";
      fips_databases — FIPS_ALL_DATABASES или список индексов баз ФИПС (0 — первая галочка).
    Ошибка по одному патенту не останавливает работу: в таблицу пишется пометка,
    и обработка идёт дальше. Возвращает RunSummary.
    """
    jobs = read_jobs(input_path)
    if not jobs:
        raise ValueError("Во входном файле нет таблиц, у которых в первой ячейке написано "
                         "«ФИПС», «Платформа» или «WIPO».")
    if not fips_databases and any(key == "fips" for key, _ in jobs):
        raise ValueError("Для поиска в ФИПС выберите хотя бы одну базу данных.")

    scrape_by_source = {
        "fips": lambda driver, request: scrapers.scrape_fips(driver, request, fips_databases),
        "platform": scrapers.scrape_platform,
        "wipo": scrapers.scrape_wipo,
    }

    report = Document(result_path) if os.path.exists(result_path) else Document()
    summary = RunSummary(result_path=result_path)

    for key, requests in jobs:
        if not requests:
            continue
        source, scrape = SOURCES[key], scrape_by_source[key]
        table = create_result_table(report)
        driver = None
        try:
            for request in requests:
                if driver is None:
                    driver = scrapers.create_driver(browser, driver_path)
                try:
                    patent = scrape(driver, request)
                except ExactResultNotFound as exc:
                    print(f"[{source.title}] {exc}")
                    add_row(table, request, NOT_FOUND_TEXT)
                    summary.not_found.append(f"{source.title}: {request}")
                except Exception as exc:  # одна ошибка не должна останавливать весь поиск
                    traceback.print_exc()
                    add_row(table, request, ERROR_TEXT.format(type(exc).__name__))
                    summary.errors.append(f"{source.title}: {request} — {_short_error(exc)}")
                    scrapers.quit_driver(driver)  # браузер мог зависнуть — следующий патент в новом
                    driver = None
                else:
                    add_row(table, patent.number, patent.name, patent.holder, patent.mpk, patent.date)
                    summary.found.append(patent)
                finally:
                    report.save(result_path)  # сохраняем после каждого патента
                    if source.browser_per_patent and driver is not None:
                        scrapers.quit_driver(driver)
                        driver = None
        finally:
            if driver is not None:
                scrapers.quit_driver(driver)

    return summary
